from calendar import month_name
from io import BytesIO
from decimal import Decimal

from django import forms
from django.contrib import messages
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction, IntegrityError
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, render, redirect
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_GET

from .expanded_withholding import MONTHS, MAX_ROWS, normalize_rows, total, import_sheet
from .models import Client, ExpandedWithholdingDraft, ExpandedWithholdingRevision
from .period_policy import validate_completed_period, default_completed_period
from .views import dashboard_permission
from .atc_catalog import ATC_ENTRIES, ATC_BY_CODE

SALT = '0619e-preview-v1'


class PeriodForm(forms.Form):
    year = forms.IntegerField(min_value=2000, max_value=9998)
    month = forms.TypedChoiceField(coerce=int, choices=[(m, month_name[m]) for m in MONTHS])

    def clean(self):
        data = super().clean()
        if data.get('year') and data.get('month'):
            try:
                validate_completed_period('MONTHLY', str(data['year']), data['month'])
            except ValidationError as exc:
                raise forms.ValidationError(exc.messages)
        return data


class RowForm(forms.Form):
    payee = forms.CharField(max_length=200)
    tin = forms.CharField(max_length=20)
    atc = forms.CharField(max_length=5)
    base = forms.DecimalField(max_digits=14, decimal_places=2, min_value=0)
    rate = forms.DecimalField(max_digits=7, decimal_places=4, min_value=0, max_value=100, label='Rate (%)', required=False, widget=forms.NumberInput(attrs={'readonly': True, 'title': 'Automatically set by the selected ATC', 'aria-readonly': 'true'}))
    reference = forms.CharField(max_length=300, required=False)

    def clean(self):
        data = super().clean()
        code = data.get('atc', '').upper()
        entry = ATC_BY_CODE.get(code)
        if code and not entry:
            self.add_error('atc', 'Select an expanded-withholding ATC from the list.')
        if entry:
            data['atc'] = code
            expected = Decimal(entry['rate'])
            if data.get('rate') is not None and data['rate'] != expected:
                self.add_error('rate', 'Rate does not match the selected ATC. Select the ATC again.')
            else:
                data['rate'] = expected
        return data


Rows = forms.formset_factory(RowForm, extra=5, can_delete=True, max_num=100, validate_max=True, absolute_max=100)


def snapshot(client):
    return {key: getattr(client, key) for key in ('registered_name', 'client_type', 'tin1', 'tin2', 'tin3', 'tin4',
                                                'rdo_code', 'registered_address', 'zip_code', 'email_address', 'version')}


@dashboard_permission('workorders.view_client', 'workorders.view_workorder')
@require_http_methods(['GET', 'POST'])
def draft(request, pk):
    client = get_object_or_404(Client, pk=pk, is_active=True)
    can_edit = request.user.has_perms(['workorders.add_workorder', 'workorders.change_workorder'])
    default = default_completed_period('MONTHLY')
    month = default['filing_month']
    if month not in MONTHS:
        month -= 1
    period = PeriodForm(request.POST if request.method == 'POST' else request.GET or
                        {'year': default['filing_year'], 'month': month})
    existing = None
    row_forms = Rows(prefix='rows')
    preview = None
    token = ''
    error = ''
    status = 200
    valid_period = period.is_valid()
    manual_allowed = True
    if valid_period:
        existing = ExpandedWithholdingDraft.objects.filter(client=client, **period.cleaned_data).first()
        manual_allowed = not existing or len(existing.rows) <= 100
        row_forms = Rows(initial=existing.rows if existing and manual_allowed else [], prefix='rows')
    if request.method == 'POST':
        if not can_edit:
            raise PermissionDenied
        action = request.POST.get('action')
        try:
            if not valid_period:
                raise ValidationError('Choose a valid completed 0619-E period.')
            version = existing.version if existing else 0
            if action == 'save':
                if request.POST.get('reviewed') != 'on':
                    raise ValidationError('Confirm the preview before saving. Preview your entries again.')
                try:
                    data = signing.loads(request.POST.get('preview_token', ''), salt=SALT, max_age=1800)
                except signing.BadSignature:
                    raise ValidationError('Preview expired or invalid. Preview your entries again.')
                if (data['client'] != str(client.pk) or data['actor'] != request.user.pk or
                    data['period'] != period.cleaned_data or data['client_version'] != client.version):
                    raise ValidationError('Client or period changed. Preview again.')
                rows = normalize_rows(data['rows'])
                with transaction.atomic():
                    if existing:
                        changed = ExpandedWithholdingDraft.objects.filter(pk=existing.pk, version=data['version']).update(
                            rows=rows, total=total(rows), version=data['version'] + 1, updated_at=timezone.now())
                        if not changed:
                            raise ValidationError('Another edit was saved. Reload and review the latest draft.')
                        existing.refresh_from_db()
                    else:
                        if data['version'] != 0:
                            raise ValidationError('Draft changed. Reload and preview again.')
                        existing = ExpandedWithholdingDraft.objects.create(client=client, **period.cleaned_data,
                                                                            rows=rows, total=total(rows))
                    ExpandedWithholdingRevision.objects.create(draft=existing, version=existing.version, rows=rows,
                        total=existing.total, actor=request.user, client_snapshot=snapshot(client))
                messages.success(request, '0619-E draft saved. No preparation or submission was queued.')
                return redirect(reverse('workorders:withholding-draft', args=[client.pk]) +
                                f'?year={existing.year}&month={existing.month}')
            if action == 'import':
                upload = request.FILES.get('workbook')
                if not upload:
                    raise ValidationError('Choose an Excel workbook.')
                preview = import_sheet(upload, request.POST.get('sheet', '').strip(), **period.cleaned_data)
            elif action == 'preview':
                if not manual_allowed:
                    raise ValidationError('Use Excel import to replace a draft with more than 100 rows.')
                row_forms = Rows(request.POST, prefix='rows')
                if not row_forms.is_valid():
                    raise ValidationError('Correct the transaction errors below.')
                preview = normalize_rows([r for r in row_forms.cleaned_data if r and not r.get('DELETE')])
                if str(version) != request.POST.get('version', '0'):
                    raise ValidationError('Another edit was saved. Reload before previewing.')
            else:
                raise ValidationError('Unknown action.')
            token = signing.dumps({'client': str(client.pk), 'actor': request.user.pk, 'period': period.cleaned_data,
                                   'client_version': client.version, 'version': version, 'rows': preview}, salt=SALT, compress=True)
        except (ValidationError, IntegrityError) as exc:
            error = ' '.join(exc.messages) if isinstance(exc, ValidationError) else 'Another draft was saved. Reload and try again.'
            preview = None
            status = 400
    return render(request, 'workorders/withholding_draft.html', {
        'taxpayer': client, 'period': period, 'row_forms': row_forms, 'draft': existing,
        'preview': preview, 'preview_total': total(preview) if preview else None, 'preview_token': token,
        'error': error, 'can_edit': can_edit, 'section': 'clients',
        'manual_allowed': manual_allowed,
        'atc_entries': ATC_ENTRIES,
        'revisions': existing.revisions.select_related('actor').order_by('-version')[:10] if existing else [],
    }, status=status)


@dashboard_permission('workorders.view_client', 'workorders.view_workorder')
@require_GET
def template(request):
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = 'Transactions'
    sheet.append(['Payee', 'TIN', 'ATC Code', 'Tax Base', 'Tax Rate', 'Withholding Tax', 'Particulars', 'Year', 'Month'])
    for row in range(2, 102):
        sheet.cell(row, 2).number_format = '@'
        sheet.cell(row, 5).number_format = '0.00%'
    stream = BytesIO()
    book.save(stream)
    response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="0619E-transactions.xlsx"'
    return response
