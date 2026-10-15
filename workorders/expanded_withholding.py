"""0619-E draft arithmetic and bounded, read-only workbook import.

Rates are supplied and reviewed by the preparer; this is not an ATC tax engine.
"""
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import re
from zipfile import ZipFile, BadZipFile

from django.core.exceptions import ValidationError

CENT = Decimal('0.01')
MONTHS = (1, 2, 4, 5, 7, 8, 10, 11)
MAX_ROWS = 500


def number(value, label, maximum, places=2):
    try:
        result = Decimal(str(value).strip().replace(',', ''))
        if not result.is_finite() or result < 0 or result > maximum:
            raise ValueError
        if result != result.quantize(Decimal(1).scaleb(-places)):
            raise ValueError
        return result
    except (InvalidOperation, ValueError):
        raise ValidationError(f'{label}: enter a nonnegative number with at most {places} decimal places.')


def normalize_rows(rows):
    if not rows or len(rows) > MAX_ROWS:
        raise ValidationError(f'Enter between 1 and {MAX_ROWS} transactions.')
    result = []
    for index, raw in enumerate(rows, 1):
        row = {key: str(raw.get(key, '') or '').strip() for key in ('payee', 'tin', 'atc', 'reference')}
        if not row['payee'] or len(row['payee']) > 200 or len(row['reference']) > 300:
            raise ValidationError(f'Row {index}: enter a payee (up to 200 characters) and reference (up to 300).')
        if not re.fullmatch(r'\d{3}-\d{3}-\d{3}-\d{3,5}', row['tin']):
            raise ValidationError(f'Row {index}: use a text TIN such as 001-002-003-00000.')
        row['atc'] = row['atc'].upper()
        if not re.fullmatch(r'W[IC]\d{3}', row['atc']):
            raise ValidationError(f'Row {index}: enter an expanded withholding ATC such as WC120; verify its applicability.')
        base = number(raw.get('base'), f'Row {index} tax base', Decimal('999999999999.99'))
        rate = number(raw.get('rate'), f'Row {index} percentage rate', Decimal('100'), 4)
        if rate == 0:
            raise ValidationError(f'Row {index}: rate must be greater than zero.')
        tax = (base * rate / 100).quantize(CENT, rounding=ROUND_HALF_UP)
        if raw.get('tax') not in (None, ''):
            supplied = number(raw['tax'], f'Row {index} withholding', Decimal('999999999999.99'))
            if supplied != tax:
                raise ValidationError(f'Row {index}: withholding does not equal tax base × rate, rounded to centavos.')
        row.update(base=f'{base:.2f}', rate=str(rate), tax=f'{tax:.2f}')
        result.append(row)
    return result


def total(rows):
    return sum((Decimal(row['tax']) for row in rows), Decimal('0.00'))


def import_sheet(upload, sheet_name, year, month):
    """Accept the provided monthly schedule or the downloadable flat template.

    Formula cells are never executed; cached values require explicit preview review.
    Multi-month Q3 schedules are deliberately rejected to avoid mixed-period imports.
    """
    from openpyxl import load_workbook
    if upload.size > 5_000_000 or not upload.name.lower().endswith('.xlsx'):
        raise ValidationError('Choose an .xlsx file no larger than 5 MB.')
    payload = upload.read()
    try:
        with ZipFile(BytesIO(payload)) as archive:
            if len(archive.infolist()) > 2000 or sum(i.file_size for i in archive.infolist()) > 30_000_000:
                raise ValidationError('Workbook is too large after decompression.')
        book = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_links=False)
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ValidationError('Unable to read this workbook. Save it as a standard .xlsx file.') from exc
    try:
        if sheet_name not in book.sheetnames:
            raise ValidationError('Sheet not found. Enter the exact monthly schedule sheet name.')
        sheet = book[sheet_name]
        if sheet.max_row > MAX_ROWS + 20:
            raise ValidationError('Sheet extends beyond the import limit. Copy the monthly rows into the template.')
        if sheet.sheet_state != 'visible':
            raise ValidationError('Choose a visible monthly schedule.')
        header = None
        rows = []
        for line, cells in enumerate(sheet.iter_rows(max_row=MAX_ROWS + 20, max_col=9, values_only=True), 1):
            values = list(cells)
            if str(values[0]).strip() in ('Registered Name', 'Payee'):
                if header is not None:
                    raise ValidationError('This sheet contains multiple schedules. Choose a single monthly sheet.')
                header = line
                expected = ['TIN', 'ATC Code', 'Tax Base', 'Tax Rate']
                if [str(v).strip() for v in values[1:5]] != expected or str(values[5]).strip() not in ('Tax', 'Withholding Tax'):
                    raise ValidationError('Columns must be Payee, TIN, ATC Code, Tax Base, Tax Rate, Withholding Tax, Particulars.')
                continue
            if header is None or not any(v is not None for v in values):
                continue
            if str(values[2]).strip().lower() == 'total':
                continue
            if len(rows) >= MAX_ROWS:
                raise ValidationError(f'Limit imports to {MAX_ROWS} transactions.')
            if values[7] is not None or values[8] is not None:
                if str(values[7]) != str(year) or str(values[8]) != str(month):
                    raise ValidationError(f'Row {line}: period does not match the selected month/year.')
            rows.append(dict(payee=values[0], tin=values[1], atc=values[2], base=values[3],
                             rate=str(Decimal(str(values[4])) * 100) if values[4] is not None else '',
                             tax=values[5], reference=values[6]))
        period = sheet['A3'].value
        if hasattr(period, 'year') and (period.year != int(year) or period.month != int(month)):
            raise ValidationError('The monthly schedule date does not match the selected period.')
        if header is None:
            raise ValidationError('No transaction header found. Use a monthly schedule or download the template.')
        return normalize_rows(rows)
    except (InvalidOperation, TypeError) as exc:
        raise ValidationError('A numeric cell is missing or invalid. Recalculate and save the workbook in Excel.') from exc
    finally:
        book.close()
